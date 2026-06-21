"""Generate printable Q_A field confirmation checklist as Word document."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml.ns import qn

OUTPUT = Path(__file__).resolve().parents[1] / "docs" / "supplementary" / "qa-field-confirmation-checklist.docx"


def set_doc_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(16)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold


def add_blank_lines(doc: Document, label: str = "", count: int = 2) -> None:
    if label:
        add_para(doc, label, bold=True)
    for _ in range(count):
        doc.add_paragraph("_" * 72)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], col_widths_cm: list[float] | None = None) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val
    if col_widths_cm:
        for row in table.rows:
            for i, w in enumerate(col_widths_cm):
                row.cells[i].width = Cm(w)


def main() -> None:
    doc = Document()
    set_doc_font(doc)

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    add_title(doc, "EDAG ARIA — Q_A 澄清清单字段确认单")
    doc.add_paragraph()

    meta = doc.add_table(rows=4, cols=2)
    meta.style = "Table Grid"
    meta.rows[0].cells[0].text = "版本"
    meta.rows[0].cells[1].text = "v1.0"
    meta.rows[1].cells[0].text = "日期"
    meta.rows[1].cells[1].text = ""
    meta.rows[2].cells[0].text = "项目名称"
    meta.rows[2].cells[1].text = ""
    meta.rows[3].cells[0].text = "填写人 / 部门 / 填写日期"
    meta.rows[3].cells[1].text = ""
    doc.add_paragraph()

    add_heading(doc, "说明", 2)
    add_para(
        doc,
        "ARIA 系统正在建设「技术澄清 Q_A 清单」模块。为确保系统生成内容与贵司实际工作方式一致，"
        "请协助确认以下事项。",
    )
    add_para(doc, "参考材料：", bold=True)
    doc.add_paragraph("• 贵司提供的 Q_A_模板.xlsx", style="List Bullet")
    doc.add_paragraph("• 前期沟通中的业务需求（待澄清问题清单、优先级、历史依据等）", style="List Bullet")
    add_para(doc, "填写完成后请回传本确认单（可勾选、补充备注）。如有更新版 Excel 模板，请一并提供。")

    add_heading(doc, "第一部分：总体方向（请勾选一项）", 1)
    add_para(doc, "系统输出的 Q_A 清单，应以哪套字段/结构为准？", bold=True)
    add_table(
        doc,
        ["选项", "说明", "勾选"],
        [
            ["A", "以 Excel 模板为准 — 严格按 Q_A_模板.xlsx 现有列（A–F 等）生成与导出", "☐"],
            ["B", "以业务需求文档为准 — 须含澄清问题、Area、优先级、历史依据等；Excel 列可扩展", "☐"],
            ["C", "两者结合 — 在现有 Excel 模板基础上，补充需求文档中的字段", "☐"],
            ["D", "其他（请说明）：", "☐"],
        ],
        [1.2, 13.5, 1.5],
    )
    add_blank_lines(doc, "备注：")

    add_heading(doc, "第二部分：字段对照确认", 1)
    add_para(doc, "请对每一行勾选「是否需要」及「由谁填写」。")
    add_para(
        doc,
        "图例：需要列 = 是否出现在导出 Excel；AI/系统 = 自动生成；EDAG = 工程师填写/修订；客户 = 客户填写。",
    )
    add_table(
        doc,
        ["#", "字段名称", "Excel\n模板", "需求\n描述", "是否需要", "主要填写方（可多选）"],
        [
            ["1", "序号（No.）", "✓ A列", "✓", "☐需要 ☐不需要", "☐AI ☐EDAG ☐客户"],
            ["2", "领域（Area）", "✓ B列", "✓", "☐需要 ☐不需要", "☐AI ☐EDAG ☐客户"],
            ["3", "提问人（Author）", "✓ C列", "—", "☐需要 ☐不需要", "☐AI ☐EDAG ☐客户"],
            ["4", "澄清问题（Question）", "✓ D列", "✓", "☐需要 ☐不需要", "☐AI ☐EDAG ☐客户"],
            ["5", "我方假设（Assumption）", "✓ E列", "△", "☐需要 ☐不需要", "☐AI ☐EDAG ☐客户"],
            ["6", "客户答复（Answer）", "✓ F列", "—", "☐需要 ☐不需要", "☐AI ☐EDAG ☐客户"],
            ["7", "影响程度/优先级（Impact）", "✗", "✓", "☐需要 ☐不需要", "☐AI ☐EDAG ☐客户"],
            ["8", "历史依据（History Reference）", "✗", "✓", "☐需要 ☐不需要", "☐AI ☐EDAG ☐客户"],
        ],
        [0.8, 3.5, 1.3, 1.3, 2.8, 4.5],
    )
    doc.add_paragraph()
    add_para(doc, "若第 7、8 项勾选「需要」，请确认 Excel 列位置与表头：", bold=True)
    add_table(
        doc,
        ["列序号", "建议表头（中文）", "建议表头（英文）", "贵司确认"],
        [
            ["第 ___ 列", "", "", "☐同意 ☐修改："],
            ["第 ___ 列", "", "", "☐同意 ☐修改："],
        ],
        [2.5, 4, 4, 5],
    )
    add_blank_lines(doc, "备注：")

    add_heading(doc, "第三部分：Excel 模板细节", 1)

    add_heading(doc, "3.1 导出给客户时，文件内容应为？", 2)
    add_table(
        doc,
        ["选项", "说明", "勾选"],
        [
            ["A", "空模板 — 仅保留表头，不含示例问题行", "☐"],
            ["B", "保留模板中的示例行 — 作为参考", "☐"],
            ["C", "仅含为本项目生成的问题 — 不含历史示例", "☐"],
            ["D", "其他：", "☐"],
        ],
        [1.2, 13.5, 1.5],
    )

    add_heading(doc, "3.2 Answer by customer（客户答复）列", 2)
    add_table(
        doc,
        ["问题", "勾选"],
        [
            ["导出给外部客户时，F 列是否应全部留空？", "☐是  ☐否"],
            ["样本模板中 F 列已有内容，是否仅为内部范例？", "☐是  ☐否"],
        ],
        [12, 4],
    )

    add_heading(doc, "3.3 Area（领域）", 2)
    add_para(doc, "请勾选官方 Area 枚举（或补充其他值）：")
    for area in [
        "Packaging", "GD&T", "Data Management", "Change Management", "BE",
        "Chassis", "EE", "CAE", "PM", "ALL / 通用",
    ]:
        doc.add_paragraph(f"☐ {area}")
    doc.add_paragraph("☐ 其他（请列出）：_______________________________________________")
    doc.add_paragraph()
    doc.add_paragraph("☐ Area 与 RFQ Function 有对应规范，请说明：_________________________")
    doc.add_paragraph("☐ 无固定规范，由工程师判断")

    add_heading(doc, "3.4 Question（问题正文）格式", 2)
    add_table(
        doc,
        ["问题", "勾选"],
        [
            ["是否要求中英双语？", "☐是 ☐否 ☐仅中文 ☐仅英文"],
            ["若双语，标准格式？", "☐英文在上/中文在下（换行） ☐其他：________"],
        ],
        [6, 10],
    )

    add_heading(doc, "3.5 Author（提问人）", 2)
    add_para(doc, "AI 或系统生成时，Author 列建议填：")
    for opt in [
        "当前登录工程师姓名/缩写",
        "固定填写「AI」或「ARIA」",
        "留空，由工程师导出前补填",
        "其他：_______________________________________________",
    ]:
        doc.add_paragraph(f"☐ {opt}")

    add_heading(doc, "3.6 其他模板问题", 2)
    add_table(
        doc,
        ["#", "问题", "贵司确认"],
        [
            ["1", "Q_A_模板.xlsx 是否为当前权威版本？", "☐是 ☐否，新版本：________"],
            ["2", "导出文件命名规范", ""],
            ["3", "是否需在标题区写入项目名、客户名、报价号？", "☐是 ☐否 位置：________"],
            ["4", "典型单次 RFQ 的 Q_A 条数范围", "约____条（最少___，最多___）"],
            ["5", "是否仅输出与当前 RFQ 范围相关的问题？", "☐是 ☐否"],
            ["6", "工程师是否需在系统内增删改后再导出？", "☐是 ☐否"],
        ],
        [0.8, 8, 7],
    )
    add_blank_lines(doc, "备注：")

    add_heading(doc, "第四部分：业务流程确认", 1)
    add_table(
        doc,
        ["#", "流程步骤", "确认"],
        [
            ["1", "EDAG 在系统内生成/编辑 Q_A 清单", "☐"],
            ["2", "导出 Excel 发给外部客户填写 Answer by customer", "☐"],
            ["3", "客户填完后回传 Excel，EDAG 内部继续评审", "☐"],
            ["4", "未来是否需要上传已填 Q_A 并读回客户答复？", "☐需要 ☐暂不需要 ☐Phase2再议"],
            ["5", "Q_A 导出前是否需与 RFQ 对标矩阵一并确认？", "☐是 ☐否"],
        ],
        [0.8, 12, 3.5],
    )
    add_blank_lines(doc, "备注：")

    add_heading(doc, "第五部分：Demo 与正式版期望", 1)
    add_table(
        doc,
        ["阶段", "贵司期望（请勾选或补充）"],
        [
            ["Demo（当前）", "☐仅需下载 Q_A Excel ☐需 AI 生成预览 ☐其他：________"],
            ["正式版（Phase 2）", "☐AI 自动填充 ☐须含历史 RAG 依据 ☐其他：________"],
        ],
        [3.5, 13],
    )
    add_blank_lines(doc, "备注：")

    add_heading(doc, "第六部分：确认签字", 1)
    add_table(
        doc,
        ["角色", "姓名", "签字", "日期"],
        [
            ["EDAG 业务确认人", "", "", ""],
            ["EDAG 技术/系统对接人", "", "", ""],
            ["ARIA 项目组", "", "", ""],
        ],
        [4.5, 3.5, 4.5, 3.5],
    )

    add_heading(doc, "附录：当前已知差异摘要（供参考）", 1)
    add_table(
        doc,
        ["项目", "Excel 模板现状", "业务需求描述现状"],
        [
            ["列数", "有效 6 列（A–F），G 列无表头", "含 Impact、History，未写入模板"],
            ["示例数据", "含约 39 条历史示例问题", "—"],
            ["Author", "模板有，需求文档未单独列出", "—"],
            ["优先级/历史依据", "模板无对应列", "需求明确要求"],
        ],
        [3, 6, 6.5],
    )

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("本确认单由 ARIA 项目组整理，用于对齐 Q_A 模块实现范围。如有疑问请联系项目组。")
    run.italic = True
    run.font.size = Pt(9)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Generated: {OUTPUT}")


if __name__ == "__main__":
    main()
