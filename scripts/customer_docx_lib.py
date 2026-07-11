"""Shared helpers for customer-facing Word documents."""

from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table


def set_cell_shading(cell, fill: str) -> None:
    from docx.oxml import OxmlElement

    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_run_font(run, size_pt: int = 10.5, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size_pt)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def new_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    return doc


def add_para(
    doc,
    text: str,
    *,
    size: int = 10.5,
    bold: bool = False,
    space_after: int = 6,
    center: bool = False,
    italic: bool = False,
):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size, bold)
    run.italic = italic
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run)


def add_numbered(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(item)
        set_run_font(run)


def add_table(doc, headers: list[str], rows: list[list[str]], *, header_fill: str = "D9E2F3") -> Table:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        r = hdr[i].paragraphs[0].add_run(h)
        set_run_font(r, 10, bold=True)
        set_cell_shading(hdr[i], header_fill)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = ""
            r = cell.paragraphs[0].add_run(val)
            set_run_font(r, 10)
    doc.add_paragraph()
    return table


def add_meta_block(doc, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        table.rows[i].cells[0].text = k
        table.rows[i].cells[1].text = v
    doc.add_paragraph()


def add_rich_para(
    doc,
    text: str,
    *,
    size: int = 10.5,
    bold: bool = False,
    italic: bool = False,
    center: bool = False,
    space_after: int = 6,
    style: str | None = None,
):
    if style:
        p = doc.add_paragraph(style=style)
    else:
        p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    add_inline_runs(p, text, size=size, bold=bold, italic=italic)
    return p


def add_inline_runs(
    paragraph,
    text: str,
    *,
    size: int = 10.5,
    bold: bool = False,
    italic: bool = False,
    mono: bool = False,
) -> None:
    import re

    pattern = re.compile(
        r"\*\*(.+?)\*\*|\*(.+?)\*|`([^`]+)`|\[([^\]]+)\]\(([^)]+)\)"
    )
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run, size, bold)
            run.italic = italic
            if mono:
                run.font.name = "Consolas"
        chunk = match.group(0)
        if chunk.startswith("**"):
            run = paragraph.add_run(match.group(1))
            set_run_font(run, size, True)
        elif chunk.startswith("*"):
            run = paragraph.add_run(match.group(2))
            set_run_font(run, size, bold)
            run.italic = True
        elif chunk.startswith("`"):
            run = paragraph.add_run(match.group(3))
            set_run_font(run, size, bold)
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        else:
            label, _url = match.group(4), match.group(5)
            run = paragraph.add_run(label)
            set_run_font(run, size, bold)
            run.italic = italic
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size, bold)
        run.italic = italic
        if mono:
            run.font.name = "Consolas"
