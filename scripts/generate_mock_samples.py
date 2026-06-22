"""Generate synthetic RFQ and mock knowledge-base documents for holiday Demo."""

from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples" / "rfq"
KB = ROOT / "backend" / "data" / "knowledge_base"

DEMO_MULTIFUNCTION_RFQ_CONTENT = """
RFQ — Demo Multi-Function Vehicle Integration

Customer: Demo OEM China
Platform: MEB Electric Vehicle Platform
Project Duration: 22 months

Scope of Work — Functions:
1. Project Management (PM)
2. Chassis — Front Suspension Design & Development
3. Chassis — Rear Suspension Design & Development
4. BIW — Body structure integration and join design
5. EE — Electrical architecture and harness routing support

Deliverables:
- M1: Concept 3D data release
- M2: Detailed 3D CAD data (Front/Rear suspension)
- M3: DMU clearance report
- M4: BIW join concept report
- M5: EE architecture interface specification

Milestones:
P1: 2026-04-01 | P2: 2026-08-01 | P3: 2026-12-01
P4: 2027-04-01 | P5: 2027-08-01 | SOP: 2027-11-01

Technical Requirements:
- Platform sharing with existing MEB derivatives
- BIW steel-aluminum mixed body structure
- EE zone architecture alignment with customer standard
- Suspension: MacPherson front, multi-link rear
"""

RFQ_CONTENT = """
RFQ — Mock Chassis Development Project

Customer: Mock Auto GmbH
Platform: MEB Electric Vehicle Platform
Project Duration: 20 months

Scope of Work — Functions:
1. Project Management (PM)
2. Chassis — Front Suspension Design & Development
3. Chassis — Rear Suspension Design & Development
4. Chassis — Steering System Integration

Deliverables:
- M1: Concept 3D data release
- M2: Detailed 3D CAD data (Front/Rear suspension)
- M3: DMU clearance report
- M4: Load decomposition report
- M5: Design validation report

Milestones:
P1: 2026-03-01 | P2: 2026-07-01 | P3: 2026-11-01
P4: 2027-03-01 | P5: 2027-07-01 | SOP: 2027-10-01

Technical Requirements:
- Platform sharing with existing MEB derivatives
- Target vehicle mass: 1850 kg
- Suspension: MacPherson front, multi-link rear
- Steering: EPS integration, rack assist type
"""

MOCK_PROJECTS = [
    {
        "dir": "mock_project_1",
        "title": "HOZON MEB Chassis Module (2023)",
        "body": """
Customer: HOZON Auto | Platform: MEB | Duration: 18 months
Functions: PM, Chassis (Front/Rear suspension, Steering)
Total manpower: ~2,400 man-days (Chassis ~1,100, PM ~180)
Key deliverables: 3D data, DMU reports, load analysis
Similarity notes: Same MEB platform, MacPherson front suspension
""",
    },
    {
        "dir": "mock_project_2",
        "title": "Mock EV Platform Chassis (2022)",
        "body": """
Customer: Mock OEM | Platform: Custom BEV | Duration: 22 months
Functions: PM, Chassis, CAE
Total manpower: ~3,100 man-days (Chassis ~1,400)
Key deliverables: Suspension design, steering integration, validation reports
Similarity notes: BEV platform, multi-link rear, EPS steering
""",
    },
    {
        "dir": "mock_project_3",
        "title": "Compact SUV Chassis Development (2021)",
        "body": """
Customer: Regional OEM | Platform: Compact SUV | Duration: 16 months
Functions: PM, Chassis
Total manpower: ~1,800 man-days (Chassis ~850)
Key deliverables: Front/rear suspension CAD, DMU, test support
Similarity notes: SUV segment, similar milestone structure P1–SOP
""",
    },
]


def write_docx(path: Path, title: str, body: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(title, level=1)
    for paragraph in body.strip().split("\n"):
        text = paragraph.strip()
        if text:
            doc.add_paragraph(text)
    doc.save(path)
    print(f"Created {path}")


def main() -> None:
    write_docx(SAMPLES / "mock_chassis_rfq.docx", "RFQ — Mock Chassis Development", RFQ_CONTENT)
    write_docx(
        SAMPLES / "demo_multifunction_rfq.docx",
        "RFQ — Demo Multi-Function Integration",
        DEMO_MULTIFUNCTION_RFQ_CONTENT,
    )

    for project in MOCK_PROJECTS:
        write_docx(
            KB / project["dir"] / "summary.docx",
            project["title"],
            project["body"],
        )

    print("Done. Synthetic samples ready for ingest.")


if __name__ == "__main__":
    main()
