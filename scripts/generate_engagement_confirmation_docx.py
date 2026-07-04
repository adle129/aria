#!/usr/bin/env python3
"""Generate Engagement package confirmation checklist for customer sign-off."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "supplementary" / "engagement-package-confirmation-checklist.docx"


def set_doc_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(16)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val


def main() -> None:
    doc = Document()
    set_doc_font(doc)
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    add_title(doc, "EDAG ARIA — Engagement 历史项目包确认单")
    doc.add_paragraph()

    meta = doc.add_table(rows=4, cols=2)
    meta.style = "Table Grid"
    meta.rows[0].cells[0].text = "版本"
    meta.rows[0].cells[1].text = "v1.2"
    meta.rows[1].cells[0].text = "日期"
    meta.rows[1].cells[1].text = date.today().isoformat()
    meta.rows[2].cells[0].text = "项目名称"
    meta.rows[2].cells[1].text = "ARIA 报价助手正式版"
    meta.rows[3].cells[0].text = "填写人 / 部门"
    meta.rows[3].cells[1].text = ""
    doc.add_paragraph()

    doc.add_heading("说明", level=2)
    add_para(
        doc,
        "R1 首期验收需要建设全类型 Engagement 知识库（不仅 RFQ）。"
        "请确认贵司将提供的脱敏历史项目包数量与每套所含文件，以便 M3 Excel 报价与 M4 QA 模块按期验收。",
    )
    add_para(doc, "建议提供 3–5 套完整项目包；至少 1 套须含历史报价 Excel 与 Q_A Excel。", bold=True)

    doc.add_heading("一、计划提供的项目包数量", level=1)
    add_table(
        doc,
        ["选项", "数量", "勾选"],
        [
            ["A", "3 套", "☐"],
            ["B", "4 套", "☐"],
            ["C", "5 套", "☐"],
            ["D", "其他：____ 套", "☐"],
        ],
    )

    doc.add_heading("二、每套项目包文件清单（请勾选）", level=1)
    add_para(doc, "以下文件类型将用于 R1 知识库入库；缺失项请备注预计提供时间。")
    add_table(
        doc,
        ["文件类型", "用途", "每套均提供", "备注"],
        [
            ["RFQ 文档（.docx / .pdf）", "RFQ 解析与对标", "☐", ""],
            ["历史 Q_A Excel", "M4 QA 行级检索", "☐", ""],
            ["历史报价人力 Excel（12 Sheet）", "M3 manpower_baselines", "☐", ""],
            ["技术方案（.pptx / .pdf）", "R1 入库；**M5 PPT 核心输入**", "☐", ""],
            ["manifest 或项目元数据（客户/车型/年份）", "Engagement 关联", "☐", ""],
        ],
    )

    doc.add_heading("三、R1 验收最低要求（双方确认）", level=1)
    add_table(
        doc,
        ["验收项", "最低标准", "确认"],
        [
            ["报价 baselines", "至少 1 份历史报价 Excel 可按 Function 查询人天", "☐"],
            ["Q_A 行检索", "至少 1 份历史 Q_A Excel 可按行检索", "☐"],
            ["历史方案检索", "至少 1 份历史技术方案可按模块/段落检索（M5 前置）", "☐"],
            ["RFQ 对标", "至少 3 份 RFQ 可完成维度确认 + 对比矩阵", "☐"],
        ],
    )

    doc.add_heading("四、项目包明细（请按行填写）", level=1)
    add_table(
        doc,
        ["序号", "项目代号/名称", "RFQ", "Q_A", "报价 Excel", "方案", "预计提供日期"],
        [
            ["1", "", "☐", "☐", "☐", "☐", ""],
            ["2", "", "☐", "☐", "☐", "☐", ""],
            ["3", "", "☐", "☐", "☐", "☐", ""],
            ["4", "", "☐", "☐", "☐", "☐", ""],
            ["5", "", "☐", "☐", "☐", "☐", ""],
        ],
    )

    doc.add_heading("五、签字确认", level=1)
    add_table(
        doc,
        ["角色", "姓名", "签字", "日期"],
        [
            ["EDAG 业务负责人", "", "", ""],
            ["EDAG IT / 数据负责人", "", "", ""],
            ["开发方项目负责人", "", "", ""],
        ],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
