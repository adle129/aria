"""Generate minimal EDAG-compatible quote_template.xlsx for Demo (replace with client template later)."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "backend" / "data" / "templates" / "quote_template.xlsx"

SHEETS = [
    "How to Use",
    "Project information",
    "Manpower",
    "PM",
    "BIW",
    "Interior",
    "GI",
    "Test validation",
    "Chassis",
    "CAE",
    "EE",
    "PS",
]

MONTH_COLUMNS = 20  # D .. W


def _setup_project_information(ws) -> None:
    labels = {
        "A6": "Customer 客户",
        "A7": "Project 项目",
        "A8": "Quotation No. 报价号",
        "A10": "Edit by 编辑人",
        "A12": "Project Start",
        "A13": "Period (Month) 周期",
        "A14": "Project End",
    }
    for cell, text in labels.items():
        ws[cell] = text
    ws["D2"] = "Milestone Timeline"
    for i in range(MONTH_COLUMNS):
        col = chr(ord("D") + i)
        ws[f"{col}2"] = f"Period {i + 1}"
        ws[f"{col}3"] = ""


def _setup_function_sheet(ws, title: str) -> None:
    ws["A1"] = title
    ws["A2"] = "Project Position"
    ws["B2"] = "Tariff Level"
    ws["C2"] = "Sum"
    for i in range(MONTH_COLUMNS):
        col = chr(ord("D") + i)
        ws[f"{col}2"] = f"Month {i + 1}"
        ws[f"{col}3"] = ""
    ws["A2"].font = Font(bold=True)
    ws["B2"].font = Font(bold=True)
    ws["C2"].font = Font(bold=True)


def _setup_manpower(ws) -> None:
    ws["A1"] = "Manpower Summary"
    ws["A2"] = "Function"
    ws["B2"] = "Description"
    ws["C2"] = "Total (man-months)"
    ws["A2"].font = Font(bold=True)
    ws["B2"].font = Font(bold=True)
    ws["C2"].font = Font(bold=True)


def generate_template(output_path: Path = OUTPUT) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    for name in SHEETS:
        ws = wb.create_sheet(name)
        if name == "How to Use":
            ws["A1"] = "EDAG Manpower Quote Template — Demo scaffold. Replace with client template."
        elif name == "Project information":
            _setup_project_information(ws)
        elif name == "Manpower":
            _setup_manpower(ws)
        elif name in {"PM", "Chassis", "BIW", "Interior", "GI", "Test validation", "CAE", "EE", "PS"}:
            _setup_function_sheet(ws, name)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    path = generate_template()
    print(f"Generated: {path}")
