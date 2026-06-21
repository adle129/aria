"""Excel manpower quote generator — EDAG 12-sheet template (Demo: PM + Chassis)."""

from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.services.generators.base import BaseGenerator
from app.services.generators.registry import GeneratorRegistry
from app.services.manpower_plan_service import MONTH_SLOTS, compute_project_end

# Project information sheet cell mapping (template-mapping.md)
PI_CUSTOMER = "B6"
PI_PROJECT = "B7"
PI_QUOTATION = "B8"
PI_EDIT_BY = "B10"
PI_START = "B12"
PI_PERIOD = "B13"
PI_END = "B14"
PI_MILESTONE_ROW = 2
PI_MILESTONE_DATE_ROW = 3
PI_MILESTONE_COL_START = 4  # column D

FUNCTION_DATA_START_ROW = 5
MONTH_COL_START = 4  # D


class ExcelManpowerGenerator(BaseGenerator):
    def generate(self, context: dict[str, Any], template_path: Path, output_path: Path) -> Path:
        if not template_path.exists():
            raise FileNotFoundError(f"Excel template not found: {template_path}")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(template_path, output_path)

        wb = load_workbook(output_path)
        rfq = context.get("rfq_modules") or {}
        plan_bundle = context.get("manpower_plan") or {}
        plan = plan_bundle.get("manpower_plan") or plan_bundle

        self._fill_project_info(
            wb["Project information"],
            rfq,
            plan_bundle.get("quotation_no"),
            context.get("edit_by", "ARIA"),
        )
        self._fill_function_sheet(wb["PM"], plan.get("PM") or [])
        self._fill_function_sheet(wb["Chassis"], plan.get("Chassis") or [])
        self._fill_manpower_summary(wb["Manpower"], plan)
        wb.save(output_path)
        return output_path

    def _fill_project_info(
        self,
        ws: Worksheet,
        rfq: dict[str, Any],
        quotation_no: str | None,
        edit_by: str,
    ) -> None:
        milestones = rfq.get("milestones") or {}
        start = milestones.get("P1")
        months = int(rfq.get("timeline_months") or 20)

        ws[PI_CUSTOMER] = rfq.get("customer", "")
        ws[PI_PROJECT] = rfq.get("project_name", "")
        ws[PI_QUOTATION] = quotation_no or ""
        ws[PI_EDIT_BY] = edit_by
        ws[PI_START] = start or ""
        ws[PI_PERIOD] = months
        ws[PI_END] = compute_project_end(start, months) or ""

        milestone_keys = ["P1", "P2", "P3", "P4", "P5", "P6", "P7", "SOP"]
        col = PI_MILESTONE_COL_START
        for key in milestone_keys:
            if key in milestones:
                ws.cell(row=PI_MILESTONE_ROW, column=col, value=key)
                ws.cell(row=PI_MILESTONE_DATE_ROW, column=col, value=milestones[key])
                col += 1

    def _fill_function_sheet(self, ws: Worksheet, rows: list[dict[str, Any]]) -> None:
        for idx, row in enumerate(rows):
            excel_row = FUNCTION_DATA_START_ROW + idx
            ws.cell(row=excel_row, column=1, value=row.get("position", ""))
            ws.cell(row=excel_row, column=2, value=row.get("tariff_level", ""))
            monthly = row.get("monthly_hours") or []
            row_sum = 0.0
            for month_idx in range(MONTH_SLOTS):
                value = monthly[month_idx] if month_idx < len(monthly) else 0
                ws.cell(row=excel_row, column=MONTH_COL_START + month_idx, value=value)
                row_sum += float(value or 0)
            ws.cell(row=excel_row, column=3, value=round(row_sum, 1))

    def _fill_manpower_summary(self, ws: Worksheet, plan: dict[str, Any]) -> None:
        summary_rows = [
            ("PM", "Project Management", plan.get("PM") or []),
            ("Chassis", "Chassis Engineering", plan.get("Chassis") or []),
        ]
        for idx, (function, desc, rows) in enumerate(summary_rows):
            excel_row = 3 + idx
            total = 0.0
            for row in rows:
                monthly = row.get("monthly_hours") or []
                total += sum(float(v or 0) for v in monthly)
            ws.cell(row=excel_row, column=1, value=function)
            ws.cell(row=excel_row, column=2, value=desc)
            ws.cell(row=excel_row, column=3, value=round(total, 1))


GeneratorRegistry.register("excel_manpower", ExcelManpowerGenerator)
