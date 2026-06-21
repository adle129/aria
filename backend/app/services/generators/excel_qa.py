"""Excel QA list export — requirements schema (5 columns, interim before customer template sign-off)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.services.generators.base import BaseGenerator
from app.services.generators.registry import GeneratorRegistry

QA_HEADERS = ["序号", "待澄清问题", "涉及功能", "影响程度", "历史依据"]
QA_KEYS = ["no", "question", "function", "impact", "history_reference"]
COLUMN_WIDTHS = [8, 48, 16, 12, 36]


def normalize_qa_item(item: dict[str, Any], index: int) -> dict[str, Any]:
    return {
        "no": item.get("no", index + 1),
        "question": item.get("question") or "",
        "function": item.get("function") or item.get("area") or "",
        "impact": item.get("impact") or "",
        "history_reference": item.get("history_reference") or "",
    }


class ExcelQAGenerator(BaseGenerator):
    def generate(self, context: dict[str, Any], template_path: Path, output_path: Path) -> Path:
        qa_items = context.get("qa_items") or []
        project_name = context.get("project_name") or ""

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "Q_A 澄清清单"

        if project_name:
            ws.cell(row=1, column=1, value=f"项目：{project_name}")
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(QA_HEADERS))
            ws.cell(row=1, column=1).font = Font(bold=True)
            header_row = 2
            data_start = 3
        else:
            header_row = 1
            data_start = 2

        for col, label in enumerate(QA_HEADERS, start=1):
            cell = ws.cell(row=header_row, column=col, value=label)
            cell.font = Font(bold=True)
            ws.column_dimensions[get_column_letter(col)].width = COLUMN_WIDTHS[col - 1]

        for idx, raw in enumerate(qa_items):
            item = normalize_qa_item(raw, idx)
            row = data_start + idx
            for col, key in enumerate(QA_KEYS, start=1):
                ws.cell(row=row, column=col, value=item[key])

        wb.save(output_path)
        return output_path


GeneratorRegistry.register("excel_qa", ExcelQAGenerator)
