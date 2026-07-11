"""Q_A Excel: one row per chunk (R1 rag-design §6)."""

from __future__ import annotations

import uuid
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# 客户 Q_A_模板.xlsx 签收版列（A–H）
QA_COLUMNS = (
    "no",
    "area",
    "author",
    "question",
    "assumption",
    "answer_by_customer",
    "impact",
    "history_reference",
)

# Labeled lines in chunk text — Area + Question first for RAG (m4-qa-merge-spec §2).
QA_CONTENT_LINES: tuple[tuple[str, str], ...] = (
    ("area", "Area"),
    ("question", "Question"),
    ("no", "No"),
    ("author", "Author"),
    ("assumption", "Assumption 我司"),
    ("answer_by_customer", "Answer by customer"),
    ("impact", "Impact"),
    ("history_reference", "History Reference"),
)


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def format_qa_row_content(payload: dict[str, Any]) -> str:
    """Structured chunk text with column labels; Area and Question lead for retrieval."""
    lines: list[str] = []
    for key, label in QA_CONTENT_LINES:
        text = _cell_str(payload.get(key))
        if key in ("area", "question", "no", "author", "assumption", "answer_by_customer"):
            lines.append(f"{label}: {text}")
        elif text:
            lines.append(f"{label}: {text}")
    return "\n".join(lines)


def load_qa_rows(path: Path) -> list[dict[str, Any]]:
    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        wb.close()
        return []

    rows: list[dict[str, Any]] = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        cells = list(row[:8]) + [None] * max(0, 8 - len(row))
        question = cells[3]
        if question is None or str(question).strip() == "":
            continue
        payload = {k: _cell_str(v) for k, v in zip(QA_COLUMNS, cells, strict=True)}
        content = format_qa_row_content(payload)
        question_text = payload["question"]
        area = payload["area"]
        rows.append(
            {
                "chunk_id": str(uuid.uuid5(uuid.NAMESPACE_URL, f"qa:{path.name}:{row_idx}")),
                "chunk_index": len(rows),
                "chunk_type": "qa_row",
                "chunk_chapter": area or f"row_{row_idx}",
                "row_number": row_idx,
                "preview": question_text[:300] + ("…" if len(question_text) > 300 else ""),
                "char_count": len(content),
                "content": content,
                "metadata": {
                    "source_doc": path.name,
                    "doc_type": "qa",
                    "row_number": row_idx,
                    **payload,
                },
            }
        )
    wb.close()
    return rows
