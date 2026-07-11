"""Quote Excel → manpower_baselines preview (Sheet-level positions, no vector index)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.services.generators.excel_manpower import FUNCTION_DATA_START_ROW, MONTH_COL_START

FUNCTION_SHEETS = (
    "PM",
    "BIW",
    "Interior",
    "GI",
    "Test validation",
    "Chassis",
    "CAE",
    "EE",
    "PS",
)

_HEADER_POSITION_LABEL = "project position"


def _is_row_hidden(ws, row: int) -> bool:
    dim = ws.row_dimensions.get(row)
    return bool(dim and dim.hidden)


def _should_skip_position_label(label: str) -> bool:
    normalized = label.strip().casefold()
    if not normalized:
        return True
    if normalized == _HEADER_POSITION_LABEL:
        return True
    if "expense" in normalized:
        return True
    if "please fill with money" in normalized:
        return True
    return False


def _position_total(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def extract_manpower_baselines(path: Path, *, include_all_positions: bool = False) -> dict[str, Any]:
    path = Path(path)
    # read_only=False: required to read row_dimensions.hidden (template placeholder rows)
    wb = load_workbook(path, read_only=False, data_only=True)
    result: dict[str, Any] = {
        "source_file": path.name,
        "project_info": {},
        "functions": {},
    }

    if "Project information" in wb.sheetnames:
        ws = wb["Project information"]
        result["project_info"] = {
            "customer": _cell(ws, "B6"),
            "project": _cell(ws, "B7"),
            "quotation_no": _cell(ws, "B8"),
            "period_months": _cell(ws, "B13"),
        }

    for sheet_name in FUNCTION_SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        positions: list[dict[str, Any]] = []
        skipped_hidden = 0
        total_man_days = 0.0
        for row in range(FUNCTION_DATA_START_ROW, (ws.max_row or 0) + 1):
            if _is_row_hidden(ws, row):
                skipped_hidden += 1
                continue
            pos = ws.cell(row=row, column=1).value
            tariff = ws.cell(row=row, column=2).value
            total = ws.cell(row=row, column=3).value
            label = "" if pos is None else str(pos).strip()
            if _should_skip_position_label(label):
                continue
            month_values = []
            for col in range(MONTH_COL_START, MONTH_COL_START + 20):
                v = ws.cell(row=row, column=col).value
                if v is not None and v != 0:
                    month_values.append({"col": col, "value": v})
            positions.append(
                {
                    "sheet": sheet_name,
                    "excel_row": row,
                    "position": label,
                    "tariff_level": "" if tariff is None else str(tariff).strip(),
                    "sum": total,
                    "nonzero_month_cells": len(month_values),
                }
            )
            total_man_days += _position_total(total)
        if positions:
            stored_positions = positions if include_all_positions else positions[:20]
            result["functions"][sheet_name] = {
                "total_man_days": round(total_man_days, 2),
                "position_count": len(positions),
                "positions": stored_positions,
                "positions_truncated": (not include_all_positions) and len(positions) > 20,
                "skipped_hidden_rows": skipped_hidden,
            }

    wb.close()
    return result


def _cell(ws, coord: str) -> Any:
    return ws[coord].value
