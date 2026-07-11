"""Unit tests for Word list numbering resolution in docx extraction."""

from __future__ import annotations

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.services.ingest.docx_body_reader import docx_to_ordered_text
from app.services.ingest.docx_numbering import DocxNumberingState, _format_counter, paragraph_text_with_numbering


def test_format_counter_chinese_and_decimal():
    assert _format_counter(1, "chineseCountingThousand") == "一"
    assert _format_counter(5, "decimal") == "5"
    assert _format_counter(2, "lowerLetter") == "b"


def _add_numbering_to_doc(doc: Document) -> None:
    """Minimal numbering.xml with abstract 12 style multi-level decimal."""
    numbering_part = doc.part.numbering_part
    root = numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "12")
    for ilvl, lvl_text in enumerate(["%1", "%1.%2"]):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(ilvl))
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal")
        lvl_text_el = OxmlElement("w:lvlText")
        lvl_text_el.set(qn("w:val"), lvl_text)
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(num_fmt)
        lvl.append(lvl_text_el)
        lvl.append(start)
        abstract.append(lvl)
    root.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), "12")
    abs_id = OxmlElement("w:abstractNumId")
    abs_id.set(qn("w:val"), "12")
    num.append(abs_id)
    root.append(num)


def _set_num_pr(paragraph, num_id: int, ilvl: int = 0) -> None:
    p_pr = paragraph._element.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl_el = OxmlElement("w:ilvl")
    ilvl_el.set(qn("w:val"), str(ilvl))
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl_el)
    num_pr.append(num_id_el)
    p_pr.append(num_pr)


def test_docx_numbering_prefixes(tmp_path):
    path = tmp_path / "numbered.docx"
    doc = Document()
    _add_numbering_to_doc(doc)
    p1 = doc.add_paragraph("First item")
    _set_num_pr(p1, 12, 1)
    p2 = doc.add_paragraph("Second item")
    _set_num_pr(p2, 12, 1)
    p3 = doc.add_paragraph("Tire size")
    _set_num_pr(p3, 12, 1)
    doc.save(path)

    loaded = Document(path)
    state = DocxNumberingState(loaded)
    texts = [paragraph_text_with_numbering(p, state) for p in loaded.paragraphs if p.text.strip()]
    assert texts[0].startswith("1.1")
    assert texts[1].startswith("1.2")
    assert texts[2].startswith("1.3")

    flat = docx_to_ordered_text(loaded)
    assert "1.1First item" in flat or "1.1 First item" in flat
    assert "1.3Tire size" in flat or "1.3 Tire size" in flat
