from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.services.generators.excel_manpower import ExcelManpowerGenerator
from app.services.manpower_plan_service import build_manpower_plan
from app.services.mock_data import MOCK_COMPARISON_TABLE, MOCK_RFQ_PARSE_RESULT

TEMPLATE = Path(__file__).resolve().parents[1] / "backend" / "data" / "templates" / "quote_template.xlsx"


@pytest.fixture
def template_path():
    if not TEMPLATE.exists():
        pytest.skip(f"Missing template: {TEMPLATE}. Run: python scripts/generate_quote_template.py")
    return TEMPLATE


def test_build_manpower_plan_has_pm_and_chassis():
    plan = build_manpower_plan(MOCK_RFQ_PARSE_RESULT, MOCK_COMPARISON_TABLE)
    assert "PM" in plan["manpower_plan"]
    assert "Chassis" in plan["manpower_plan"]
    assert len(plan["manpower_plan"]["PM"]) >= 1
    assert len(plan["manpower_plan"]["Chassis"]) >= 1


def test_generate_quote_creates_file(template_path, tmp_path):
    plan = build_manpower_plan(MOCK_RFQ_PARSE_RESULT, MOCK_COMPARISON_TABLE)
    output = tmp_path / "quote_test.xlsx"
    gen = ExcelManpowerGenerator()
    gen.generate(
        {"rfq_modules": MOCK_RFQ_PARSE_RESULT, "manpower_plan": plan},
        template_path,
        output,
    )
    assert output.exists()
    wb = load_workbook(output)
    assert "PM" in wb.sheetnames
    assert "Chassis" in wb.sheetnames
    assert wb["Project information"]["B6"].value == "Mock Auto GmbH"
    assert wb["Project information"]["B7"].value == "Mock Chassis Development"
    pm_sum = wb["PM"]["C5"].value
    assert pm_sum and float(pm_sum) > 0


def test_generate_quote_empty_modules(template_path, tmp_path):
    plan = {"manpower_plan": {"PM": [], "Chassis": []}, "quotation_no": "260001B000"}
    output = tmp_path / "quote_empty.xlsx"
    gen = ExcelManpowerGenerator()
    gen.generate(
        {"rfq_modules": {"project_name": "Empty", "customer": "X", "timeline_months": 12}, "manpower_plan": plan},
        template_path,
        output,
    )
    assert output.exists()
    wb = load_workbook(output)
    assert wb["Manpower"]["C3"].value == 0
