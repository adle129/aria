"""Unit tests for RFQ document loader (.docx / .doc)."""

from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock

import pytest

from app.services.ingest.rfq_document_loader import load_rfq_text


def test_load_docx(tmp_path):
    from docx import Document

    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("3.1 Project scope")
    doc.add_paragraph("Chassis development")
    doc.save(path)

    text, loader = load_rfq_text(path)
    assert "Project scope" in text
    assert loader == "python-docx"


def test_load_doc_linux_uses_libreoffice(monkeypatch, tmp_path):
    doc_path = tmp_path / "legacy.doc"
    doc_path.write_bytes(b"fake-doc")

    monkeypatch.setattr("app.services.ingest.rfq_document_loader.sys.platform", "linux")
    monkeypatch.setattr("app.services.ingest.rfq_document_loader._find_soffice", lambda: "/usr/bin/soffice")

    converted = tmp_path / "converted.docx"
    from docx import Document

    out = Document()
    out.add_paragraph("RFQ section 3.2")
    out.save(converted)

    def fake_run(cmd, **kwargs):
        outdir = Path(cmd[cmd.index("--outdir") + 1])
        target = outdir / "legacy.docx"
        target.write_bytes(converted.read_bytes())
        return MagicMock(returncode=0, stdout="", stderr="")

    monkeypatch.setattr("app.services.ingest.rfq_document_loader.subprocess.run", fake_run)

    text, loader = load_rfq_text(doc_path)
    assert "RFQ section 3.2" in text
    assert loader == "libreoffice+python-docx"


def test_load_docx_interleaves_tables(tmp_path):
    from docx import Document

    path = tmp_path / "ordered.docx"
    doc = Document()
    doc.add_paragraph("1.4 Tire size")
    doc.add_paragraph("1.5 Key systems:")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Range"
    table.cell(1, 1).text = "600km"
    doc.add_paragraph("Section two")
    doc.save(path)

    text, loader = load_rfq_text(path)
    assert loader == "python-docx"
    idx_header = text.find("1.5 Key systems")
    idx_table = text.find("[TABLE] Item | Value")
    idx_section2 = text.find("Section two")
    assert idx_header >= 0 and idx_table > idx_header and idx_section2 > idx_table


def test_load_doc_linux_without_libreoffice(monkeypatch, tmp_path):
    doc_path = tmp_path / "legacy.doc"
    doc_path.write_bytes(b"fake-doc")
    monkeypatch.setattr("app.services.ingest.rfq_document_loader.sys.platform", "linux")
    monkeypatch.setattr("app.services.ingest.rfq_document_loader._find_soffice", lambda: None)

    with pytest.raises(ValueError, match="LibreOffice"):
        load_rfq_text(doc_path)
