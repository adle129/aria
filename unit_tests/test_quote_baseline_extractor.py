"""Unit tests for quote Excel baseline extraction."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from app.services.generators.excel_manpower import FUNCTION_DATA_START_ROW
from app.services.ingest.quote_baseline_extractor import extract_manpower_baselines


def _write_ps_sheet(path: Path, *, hide_row: int | None = None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "PS"
    ws.cell(row=5, column=1, value="PS (Headcount Internal& External）")
    ws.cell(row=72, column=1, value="PS External Expense (Please fill with Money）")
    ws.cell(row=73, column=1, value="XXXX")
    ws.cell(row=87, column=1, value="PS Travel Expense (Please fill with Money）")
    if hide_row is not None:
        ws.row_dimensions[hide_row].hidden = True
    wb.save(path)
    wb.close()


def test_extract_skips_hidden_placeholder_row(tmp_path):
    xlsx = tmp_path / "quote.xlsx"
    _write_ps_sheet(xlsx, hide_row=73)

    result = extract_manpower_baselines(xlsx)
    ps = result["functions"]["PS"]
    labels = [p["position"] for p in ps["positions"]]

    assert ps["position_count"] == 1
    assert "XXXX" not in labels
    assert "Expense" not in labels[0]
    assert ps["skipped_hidden_rows"] >= 1
    assert ps["positions"][0]["excel_row"] == 5
    assert ps["positions"][0]["sheet"] == "PS"


def test_extract_skips_expense_and_money_rows(tmp_path):
    xlsx = tmp_path / "quote_expense.xlsx"
    _write_ps_sheet(xlsx, hide_row=73)

    result = extract_manpower_baselines(xlsx)
    labels = [p["position"] for p in result["functions"]["PS"]["positions"]]
    assert len(labels) == 1
    assert all("expense" not in label.casefold() for label in labels)
    assert "XXXX" not in labels


def test_extract_includes_visible_placeholder_row(tmp_path):
    xlsx = tmp_path / "quote_visible.xlsx"
    _write_ps_sheet(xlsx, hide_row=None)

    result = extract_manpower_baselines(xlsx)
    labels = [p["position"] for p in result["functions"]["PS"]["positions"]]
    assert "XXXX" in labels


@pytest.mark.skipif(
    not Path(r"E:/AI文档项目/RE_ 报价AI需求沟通/报价人力模板.xlsx").exists(),
    reason="validation corpus not on this machine",
)
def test_validation_corpus_ps_has_headcount_rows_without_expense():
    path = Path(r"E:/AI文档项目/RE_ 报价AI需求沟通/报价人力模板.xlsx")
    result = extract_manpower_baselines(path)
    ps = result["functions"]["PS"]
    labels = [p["position"] for p in ps["positions"]]
    assert ps["position_count"] >= 1
    assert "XXXX" not in labels
    assert all("expense" not in label.casefold() for label in labels)
