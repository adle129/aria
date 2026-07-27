import json
from pathlib import Path
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook

from app.config import Settings
from app.services.engagement_ingest_service import EngagementIngestError, EngagementIngestService
from app.services.ingest.chunk_benchmarks import assert_vector_chunks_rfqa_only
from app.services.knowledge_index_service import flatten_engagement_chunks


SAMPLE_RFQ = Path(__file__).resolve().parents[1] / "samples" / "rfq" / "mock_chassis_rfq.docx"
QA_TEMPLATE = Path(__file__).resolve().parents[1] / "backend" / "data" / "templates" / "qa_template.xlsx"


@pytest.fixture
def engagement_folder(tmp_path):
    if not SAMPLE_RFQ.exists() or not QA_TEMPLATE.exists():
        pytest.skip("sample rfq or qa template missing")
    folder = tmp_path / "test_engagement"
    folder.mkdir()
    (folder / "RFQ_mock.docx").write_bytes(SAMPLE_RFQ.read_bytes())
    (folder / "Q_A_mock.xlsx").write_bytes(QA_TEMPLATE.read_bytes())
    manifest = {
        "engagement_id": "test_engagement",
        "project_name": "Test Engagement",
        "customer": "OEM-Test",
        "year": 2026,
        "functions": ["Chassis", "PM"],
        "documents": [
            {"path": "RFQ_mock.docx", "doc_type": "rfq"},
            {"path": "Q_A_mock.xlsx", "doc_type": "qa"},
        ],
    }
    (folder / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    return folder


def test_assert_rfq_gate_accepts_rfq_only():
    EngagementIngestService.assert_rfq_gate(
        [{"metadata": {"doc_type": "rfq"}}],
        "x",
    )


def test_assert_rfq_gate_rejects_engagement_without_rfq():
    with pytest.raises(EngagementIngestError):
        EngagementIngestService.assert_rfq_gate(
            [{"metadata": {"doc_type": "qa"}}],
            "x",
        )


def test_assert_metadata_gate_rejects_incomplete():
    from app.schemas.engagement import EngagementManifest

    with pytest.raises(EngagementIngestError, match="项目信息未齐"):
        EngagementIngestService.assert_metadata_gate(
            EngagementManifest(
                engagement_id="x",
                project_name="Only Name",
                documents=[],
            )
        )


def test_import_all_skips_incomplete_metadata(tmp_path, monkeypatch):
    if not SAMPLE_RFQ.exists():
        pytest.skip("sample rfq missing")
    kb = tmp_path / "kb"
    incomplete = kb / "incomplete_meta"
    incomplete.mkdir(parents=True)
    (incomplete / "RFQ_mock.docx").write_bytes(SAMPLE_RFQ.read_bytes())
    (incomplete / "manifest.json").write_text(
        json.dumps(
            {
                "engagement_id": "incomplete_meta",
                "project_name": "Incomplete",
                "documents": [{"path": "RFQ_mock.docx", "doc_type": "rfq"}],
            }
        ),
        encoding="utf-8",
    )
    complete = kb / "complete_meta"
    complete.mkdir(parents=True)
    (complete / "RFQ_mock.docx").write_bytes(SAMPLE_RFQ.read_bytes())
    (complete / "manifest.json").write_text(
        json.dumps(
            {
                "engagement_id": "complete_meta",
                "project_name": "Complete",
                "customer": "OEM-A",
                "year": 2024,
                "functions": ["PM"],
                "documents": [{"path": "RFQ_mock.docx", "doc_type": "rfq"}],
            }
        ),
        encoding="utf-8",
    )

    settings = Settings(
        knowledge_base_path=str(kb),
        mock_rag=False,
        manpower_baselines_path=str(tmp_path / "baselines.json"),
        database_url="sqlite://",
    )
    service = EngagementIngestService(settings, db=None)
    seen_engagement_ids: list[str] = []

    class FakeIndex:
        def __init__(self, *_args, **_kwargs):
            pass

        def build_generation(
            self,
            chunks,
            *,
            corpus_path=None,
            source_file=None,
            created_by_job_id=None,
            request_type="kb_full",
        ):
            for c in chunks:
                eid = (c.get("metadata") or {}).get("engagement_id")
                if eid and eid not in seen_engagement_ids:
                    seen_engagement_ids.append(eid)
            return {"generation_id": "staged-meta", "chunk_count": len(chunks)}

        def activate_generation(self, staged):
            return {
                **staged,
                "active_generation": staged["generation_id"],
                "last_index_at": "2026-07-26T00:00:00Z",
            }

        def fail_generation(self, generation_id, error):
            return None

    monkeypatch.setattr(
        "app.services.engagement_ingest_service.KnowledgeIndexService",
        FakeIndex,
    )
    result = service.import_all()
    assert "incomplete_meta" not in seen_engagement_ids
    assert "complete_meta" in seen_engagement_ids
    failed_paths = {f["path"] for f in result["failed_files"]}
    assert "incomplete_meta" in failed_paths
    by_id = {e["engagement_id"]: e for e in result["engagements"]}
    assert by_id["incomplete_meta"]["status"] == "failed"
    assert "metadata" in by_id["incomplete_meta"]["missing"]
    assert by_id["complete_meta"]["status"] == "indexed"


def test_prepare_engagement_builds_rfqa_chunks(engagement_folder, tmp_path):
    kb = tmp_path / "kb"
    kb.mkdir()
    target = kb / engagement_folder.name
    target.mkdir()
    for item in engagement_folder.iterdir():
        target.joinpath(item.name).write_bytes(item.read_bytes())

    service = EngagementIngestService(
        Settings(
            knowledge_base_path=str(kb),
            mock_rag=False,
            manpower_baselines_path=str(tmp_path / "baselines.json"),
        )
    )
    manifest, chunks, baseline = service.prepare_engagement(target)
    assert manifest.engagement_id == "test_engagement"
    doc_types = {(c["metadata"] or {}).get("doc_type") for c in chunks}
    assert "rfq" in doc_types
    assert "qa" in doc_types
    assert_vector_chunks_rfqa_only(chunks)
    assert all(c["metadata"]["engagement_id"] == "test_engagement" for c in chunks)
    assert baseline is None


def test_prepare_engagement_accepts_rfq_only(tmp_path):
    if not SAMPLE_RFQ.exists():
        pytest.skip("sample rfq missing")
    kb = tmp_path / "kb"
    target = kb / "copper_engagement"
    target.mkdir(parents=True)
    (target / "RFQ_mock.docx").write_bytes(SAMPLE_RFQ.read_bytes())
    (target / "manifest.json").write_text(
        json.dumps(
            {
                "engagement_id": "copper_engagement",
                "project_name": "Copper Engagement",
                "customer": "OEM-Copper",
                "year": 2025,
                "functions": ["Chassis"],
                "documents": [{"path": "RFQ_mock.docx", "doc_type": "rfq"}],
            }
        ),
        encoding="utf-8",
    )
    service = EngagementIngestService(
        Settings(
            knowledge_base_path=str(kb),
            mock_rag=False,
            manpower_baselines_path=str(tmp_path / "baselines.json"),
        )
    )

    manifest, chunks, baseline = service.prepare_engagement(target)

    assert manifest.engagement_id == "copper_engagement"
    assert {(c["metadata"] or {}).get("doc_type") for c in chunks} == {"rfq"}
    assert baseline is None


def test_import_all_mock_embed(engagement_folder, tmp_path, monkeypatch):
    kb = tmp_path / "kb"
    kb.mkdir()
    target = kb / engagement_folder.name
    target.mkdir()
    for item in engagement_folder.iterdir():
        target.joinpath(item.name).write_bytes(item.read_bytes())

    settings = Settings(
        knowledge_base_path=str(kb),
        mock_rag=False,
        manpower_baselines_path=str(tmp_path / "baselines.json"),
        database_url="sqlite://",
    )
    service = EngagementIngestService(settings, db=None)

    class FakeIndex:
        def __init__(self, *_args, **_kwargs):
            pass

        def build_generation(
            self,
            chunks,
            *,
            corpus_path=None,
            source_file=None,
            created_by_job_id=None,
            request_type="kb_full",
        ):
            rfq = sum(1 for c in chunks if (c.get("metadata") or {}).get("doc_type") == "rfq")
            qa = sum(1 for c in chunks if (c.get("metadata") or {}).get("doc_type") == "qa")
            assert rfq >= 1 and qa >= 1
            return {"generation_id": "staged-1", "chunk_count": len(chunks)}

        def activate_generation(self, staged):
            return {
                **staged,
                "active_generation": staged["generation_id"],
                "last_index_at": "2026-07-07T00:00:00Z",
            }

        def fail_generation(self, generation_id, error):
            return None

    monkeypatch.setattr(
        "app.services.engagement_ingest_service.KnowledgeIndexService",
        FakeIndex,
    )
    result = service.import_all()
    assert result["new_chunks"] >= 1
    assert result["doc_type_counts"]["rfq"] >= 1
    assert result["doc_type_counts"]["qa"] >= 1
    assert result["engagements"][0]["status"] == "indexed"
    assert result["engagements"][0]["tier"] == "silver"
    assert result["engagements"][0]["indexable"] is True


def _write_min_quote(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "PM"
    ws.cell(row=5, column=1, value="Project Manager")
    ws.cell(row=3, column=3, value=10)
    ws.cell(row=72, column=1, value="PM Travel Expense (Please fill with Money）")
    wb.save(path)
    wb.close()


def test_prepare_engagement_extracts_baselines_not_vectors(engagement_folder, tmp_path):
    kb = tmp_path / "kb"
    kb.mkdir()
    target = kb / engagement_folder.name
    target.mkdir()
    for item in engagement_folder.iterdir():
        target.joinpath(item.name).write_bytes(item.read_bytes())
    quote_path = target / "quote_manpower.xlsx"
    _write_min_quote(quote_path)
    manifest_path = target / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["documents"].append({"path": "quote_manpower.xlsx", "doc_type": "quote_manpower"})
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

    baselines_path = tmp_path / "baselines.json"
    service = EngagementIngestService(
        Settings(
            knowledge_base_path=str(kb),
            mock_rag=False,
            manpower_baselines_path=str(baselines_path),
        )
    )
    _manifest, chunks, baseline = service.prepare_engagement(target)
    assert_vector_chunks_rfqa_only(chunks)
    assert baseline is not None
    assert "PM" in baseline["functions"]
    pm_positions = baseline["functions"]["PM"]["positions"]
    assert all("expense" not in p["position"].casefold() for p in pm_positions)
